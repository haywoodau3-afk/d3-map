from __future__ import annotations

import csv
import hashlib
import json
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from .errors import DMapError, ValidationError


def _dependencies() -> tuple[Any, ...]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            Image,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as error:
        raise DMapError(
            "PDF and XLSX reports require the standard d3map report dependencies; reinstall d3map-catalysis"
        ) from error
    return (
        Workbook,
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
        colors,
        TA_CENTER,
        letter,
        ParagraphStyle,
        getSampleStyleSheet,
        inch,
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValidationError(f"expected an object in {path}")
    return value


def _csv_rows(path: Path) -> list[list[Any]]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8", newline="") as stream:
        return [list(row) for row in csv.reader(stream)]


def _flatten(value: Any, prefix: str = "") -> Iterable[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            yield from _flatten(child, path)
    elif isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            yield prefix, json.dumps(value, ensure_ascii=False)
        else:
            for index, child in enumerate(value):
                yield from _flatten(child, f"{prefix}[{index}]")
    else:
        yield prefix, value


def _style_workbook(workbook: Any) -> None:
    _, Alignment, Border, Font, PatternFill, Side, *_ = _dependencies()
    header_fill = PatternFill("solid", fgColor="183153")
    stripe_fill = PatternFill("solid", fgColor="EEF3F8")
    border = Border(bottom=Side(style="thin", color="D9D9D9"))
    for index, worksheet in enumerate(workbook.worksheets):
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2" if worksheet.max_row > 8 else None
        worksheet.sheet_properties.tabColor = "183153" if index == 0 else "6E8EAD"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_index, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="202833")
                cell.alignment = Alignment(vertical="center")
                cell.border = border
                if row_index % 2 == 0:
                    cell.fill = stripe_fill
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for column in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column]
            width = min(max(max((len(value) for value in values), default=0) + 2, 11), 48)
            worksheet.column_dimensions[column[0].column_letter].width = width
        if worksheet.max_row > 1 and worksheet.max_column > 1:
            worksheet.auto_filter.ref = worksheet.dimensions


def _write_rows(worksheet: Any, rows: list[list[Any]]) -> None:
    for row in rows:
        worksheet.append(row)


def _summary_rows(descriptors: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = [["Descriptor", "Statistic", "Value", "Unit"]]
    for descriptor, summary in descriptors["ensemble_scalar_descriptors"].items():
        unit = "%"
        for statistic in (
            "mean", "population_standard_deviation", "percentile_5", "percentile_95",
            "percentile_span", "minimum", "maximum", "lowest_energy_value",
        ):
            rows.append([descriptor, statistic, summary.get(statistic), unit])
    rows.extend(
        [
            ["ensemble", "analysis_mode", descriptors.get("analysis_mode"), ""],
            ["ensemble", "population_model", descriptors.get("population_model"), ""],
            ["ensemble", "temperature", descriptors.get("temperature"), "K"],
            ["ensemble", "effective_ensemble_size", descriptors.get("effective_ensemble_size"), ""],
        ]
    )
    return rows


def _workbook_for_structure(output: Path, role: str) -> Path:
    Workbook, *_ = _dependencies()
    descriptors = _read_json(output / "descriptors.json")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_rows(summary, _summary_rows(descriptors))
    provenance = workbook.create_sheet("Provenance")
    _write_rows(
        provenance,
        [["Field", "Value"]]
        + [[key, value] for key, value in _flatten(descriptors) if not key.startswith("features")],
    )
    baseline_path = output / "static-baselines.json"
    if baseline_path.is_file():
        baselines = _read_json(baseline_path)["baselines"]
        worksheet = workbook.create_sheet("Static baselines")
        _write_rows(
            worksheet,
            [["Baseline", "Status", "Definition or reason", "Vbur (%)", "G (%)"]]
            + [
                [
                    name,
                    values.get("status"),
                    values.get("definition", values.get("reason")),
                    values.get("vbur_percent"),
                    values.get("g_percent"),
                ]
                for name, values in baselines.items()
            ],
        )
    sources = (
        ("Per geometry", output / "plot-data/per-geometry-descriptors.csv"),
        ("Axial profile", output / "axial-profile.csv"),
        ("Displaced scan", output / "plot-data/displaced-scan.csv"),
        ("Features", output / "features.csv"),
        ("ML spatial", output / "plot-data/ml-spatial-fractions.csv"),
        ("ML direction", output / "plot-data/ml-direction-offset-features.csv"),
    )
    for name, path in sources:
        rows = _csv_rows(path)
        if rows:
            worksheet = workbook.create_sheet(name)
            _write_rows(worksheet, rows)
    dictionary = workbook.create_sheet("Data dictionary")
    _write_rows(
        dictionary,
        [
            ["Term", "Definition"],
            ["Role", role],
            ["Vbur", "SambVca 2.1-compatible buried volume percentage"],
            ["G", "Guzei-Wendt method-equivalent solid-angle percentage"],
            ["Population", "The explicitly labelled ensemble or trajectory weight"],
            ["Difference", "Extended minus reference in d3-map outputs"],
            ["Missing value", "Unavailable or outside-domain value; it is not numerical zero"],
        ],
    )
    _style_workbook(workbook)
    target = output / "d3map-results.xlsx"
    workbook.save(target)
    return target


def _workbook_for_comparison(output: Path) -> Path:
    Workbook, *_ = _dependencies()
    comparison = _read_json(output / "comparison.json")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Comparison"
    rows = _csv_rows(output / "comparison.csv")
    _write_rows(summary, rows)
    baseline_rows = comparison.get("static_baseline_comparisons", [])
    if baseline_rows:
        baseline = workbook.create_sheet("Static baselines")
        headings = list(baseline_rows[0])
        _write_rows(
            baseline,
            [headings] + [[row.get(key) for key in headings] for row in baseline_rows],
        )
    settings = workbook.create_sheet("Settings")
    _write_rows(
        settings,
        [["Field", "Value"]]
        + [[key, value] for key, value in _flatten(comparison) if key != "scalar_comparisons"],
    )
    dictionary = workbook.create_sheet("Data dictionary")
    _write_rows(
        dictionary,
        [
            ["Term", "Definition"],
            ["Reference", "The subtraction baseline"],
            ["Extended", "The independently sampled structure being compared"],
            ["Signed difference", "Extended minus reference"],
            ["Absolute difference", "Magnitude of the signed difference"],
            ["Relative difference", "Signed difference divided by reference when safely nonzero"],
            ["Tolerance mask", "A visualization mask that never changes raw numerical differences"],
        ],
    )
    _style_workbook(workbook)
    target = output / "d3map-comparison.xlsx"
    workbook.save(target)
    return target


def _pdf_document(output: Path, *, title: str, paragraphs: list[str], table_rows: list[list[Any]], image: Path | None) -> Path:
    (
        _Workbook,
        _Alignment,
        _Border,
        _Font,
        _PatternFill,
        _Side,
        colors,
        TA_CENTER,
        letter,
        ParagraphStyle,
        getSampleStyleSheet,
        inch,
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    ) = _dependencies()
    target = output / "d3map-summary.pdf"
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="D3Title",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#18202A"),
            alignment=TA_CENTER,
            spaceAfter=18,
        )
    )
    document = SimpleDocTemplate(
        str(target), pagesize=letter, rightMargin=0.65 * inch, leftMargin=0.65 * inch,
        topMargin=0.65 * inch, bottomMargin=0.65 * inch,
        title=title, author="d3map",
    )
    story: list[Any] = [Paragraph(title, styles["D3Title"])]
    for text in paragraphs:
        story.extend([Paragraph(text, styles["BodyText"]), Spacer(1, 8)])
    if table_rows:
        printable = [
            [
                ""
                if value is None
                else f"{value:.4f}"
                if isinstance(value, float)
                else str(value)
                for value in row
            ]
            for row in table_rows
        ]
        table = Table(printable, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183153")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(table)
    if image is not None and image.is_file():
        story.extend([PageBreak(), Paragraph("Principal result figure", styles["Heading1"]), Spacer(1, 8)])
        rendered = Image(str(image))
        rendered._restrictSize(7.0 * inch, 8.7 * inch)
        story.append(rendered)

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#5D6875"))
        canvas.drawString(0.65 * inch, 0.35 * inch, "Generated locally by d3map")
        canvas.drawRightString(7.85 * inch, 0.35 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return target


def write_structure_reports(output: str | Path, *, role: str) -> tuple[Path, Path]:
    directory = Path(output).resolve()
    descriptors = _read_json(directory / "descriptors.json")
    workbook = _workbook_for_structure(directory, role)
    summary = _summary_rows(descriptors)
    temperature = descriptors.get("temperature")
    temperature_text = "not provided" if temperature is None else f"{temperature} K"
    paragraphs = [
        f"This {role} d2-map report summarizes the static and ensemble-resolved steric calculation.",
        "Buried volume is reported using the SambVca 2.1-compatible profile. Solid angle is reported using the Guzei-Wendt method-equivalent profile.",
        f"Analysis mode: {descriptors.get('analysis_mode')}. Population model: {descriptors.get('population_model')}. Temperature: {temperature_text}.",
        "The complete interactive maps, numerical fields, aligned structures, CSV and JSON data remain in the accompanying result directory.",
    ]
    image = directory / "figures/all-features-dashboard.png"
    pdf = _pdf_document(directory, title=f"d3map {role} summary", paragraphs=paragraphs, table_rows=summary[:17], image=image)
    return pdf, workbook


def write_comparison_reports(output: str | Path) -> tuple[Path, Path]:
    directory = Path(output).resolve()
    comparison = _read_json(directory / "comparison.json")
    workbook = _workbook_for_comparison(directory)
    rows = _csv_rows(directory / "comparison.csv")
    paragraphs = [
        "This d3-map report compares independently sampled reference and extended structures.",
        "Every signed difference is extended minus reference. Tolerance masking applies only to visual classification and never changes the stored numerical difference fields.",
        f"Comparison status: {comparison.get('status')}. Compatibility mismatches: {comparison.get('compatibility_mismatches') or 'none'}.",
        "Complete d2-map reports for both structures and machine-readable paired fields accompany this summary.",
    ]
    pdf = _pdf_document(directory, title="d3map paired comparison summary", paragraphs=paragraphs, table_rows=rows[:18], image=None)
    return pdf, workbook


def refresh_manifest(directory: str | Path) -> Path:
    root = Path(directory).resolve()
    path = root / "manifest.json"
    payload = _read_json(path) if path.is_file() else {}
    payload.update(
        {
            "artifacts": {
                str(item.relative_to(root)): hashlib.sha256(item.read_bytes()).hexdigest()
                for item in sorted(root.rglob("*"))
                if item.is_file() and item.name != "manifest.json"
            }
        }
    )
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return path
